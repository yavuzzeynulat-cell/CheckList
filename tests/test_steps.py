"""Ekranlar: dosya birak -> filtre -> sonuc; Settings'te kural ekle/kapat."""
import openpyxl

import main
import output
import settings
import shell
import steps
from fixtures import CHECK_HEADER, LAB_HEADER, RIA_HEADER, check, ria, xlsx

A = "26437-RIA-04C-QC-PIT-OP-00004"
B = "26437-RIA-02B-DR-PR2-BX-00031"
C = "26437-RIA-04D-EW-REX-MW-00906"


def _shell(tk_root, tmp_path):
    settings.save_out_dir(tmp_path / "out")
    (tmp_path / "out").mkdir()
    sh = shell.Shell(tk_root, steps.all_steps(), version="0.1.0")
    sh.steps[2].opener = lambda p: None
    return sh


def _files(tmp_path):
    return [xlsx(tmp_path / "DataGrid.xlsx", RIA_HEADER,
                 [ria(A), ria(B), ria(C), ria("26437-RIA-02B-ST-IN2-RW-00571", title="Reinforcement of wall")]),
            xlsx(tmp_path / "DataGrid (1).xlsx", CHECK_HEADER, [check("CHL-1", A)]),
            xlsx(tmp_path / "DataGrid (2).xlsx", LAB_HEADER, [["x", "y", "z"]])]


def test_full_flow(tk_root, tmp_path):
    sh = _shell(tk_root, tmp_path)
    files, flt, res = sh.steps[0], sh.steps[1], sh.steps[2]
    assert not files.next_btn.enabled
    files.add(_files(tmp_path), wait=True)
    tk_root.update()
    s = sh.ctx.session
    assert (len(s.rias), len(s.checklists), len(s.others)) == (1, 1, 1)
    assert files.next_btn.enabled
    sh.select("filter")
    assert flt.kept() == (3, 4)          # REX varsayilan kapali yok (testte exclude {}), Reinforcement kurali eler
    flt.toggle("Activity", "PR2")
    assert flt.kept() == (2, 4)
    assert settings.load_exclude()["Activity"] >= {"PR2"}
    flt.compare()
    assert sh.active is res
    wb = openpyxl.load_workbook(s.out_path)
    assert wb[output.MATCHED_SHEET].cell(2, 2).value == A
    assert [r[1].value for r in wb[output.MISSING_SHEET].iter_rows(min_row=2)] == [C]
    res.set_view("has")
    res.start_over()
    assert sh.active is files and not s.tables


def test_add_toggle_delete_rule(tk_root, tmp_path):
    sh = _shell(tk_root, tmp_path)
    prefs = sh.steps[3]
    sh.select("settings")
    assert prefs.add_rule("Document Title", "contains", "Rebar")
    assert not prefs.add_rule("Document Title", "contains", "rebar")     # ayni kural
    assert not prefs.add_rule("", "starts with", "x")
    assert [r["text"] for r in settings.load_rules()] == ["Reinforcement", "Rebar"]
    prefs.toggle_rule(0)
    assert settings.load_rules()[0]["on"] is False
    prefs.delete_rule(1)
    assert [r["text"] for r in settings.load_rules()] == ["Reinforcement"]


def test_reset_needs_two_clicks(tk_root, tmp_path):
    sh = _shell(tk_root, tmp_path)
    prefs = sh.steps[3]
    settings.save_rules([])
    assert prefs.reset_all() is False
    assert prefs.reset_all() is True
    assert settings.load_rules()[0]["text"] == "Reinforcement"


def test_default_filter_add_remove(tk_root, tmp_path):
    sh = _shell(tk_root, tmp_path)
    prefs = sh.steps[3]
    settings.save_default_exclude({"Activity": {"REX"}})
    assert prefs.add_default("Discipline", " qc ")            # buyuk harfe, bosluksuz
    assert not prefs.add_default("Discipline", "QC")          # zaten var
    assert not prefs.add_default("Discipline", "")
    assert not prefs.add_default("Nope", "X")                 # oyle parca yok
    assert settings.load_default_exclude() == {"Activity": {"REX"}, "Discipline": {"QC"}}
    prefs.remove_default("Activity", "REX")
    assert settings.load_default_exclude() == {"Discipline": {"QC"}}
    settings.reset_exclude()
    assert settings.load_exclude() == {"Discipline": {"QC"}}  # Reset bu secimi getirir


def test_default_filter_dialog(tk_root, tmp_path):
    sh = _shell(tk_root, tmp_path)
    d = sh.steps[3].edit_default()
    d.set_all("Activity", True)
    d.save()
    assert settings.load_default_exclude() == {}


def test_drop_still_accepted_after_window_fades_in(tk_root, tmp_path):
    """Tk'nin -alpha ayari WS_EX_ACCEPTFILES'i siliyordu (olculdu 2026-10-14):
    pencere solarak acildiktan sonra Windows birakilan dosyayi hic gondermiyordu."""
    import ctypes
    sh = _shell(tk_root, tmp_path)
    sh.show()
    for _ in range(40):                 # acilis animasyonu (~220 ms) bitsin
        tk_root.update()
        tk_root.after(15)
    tk_root.update()
    hook = sh.steps[0]._hook
    assert hook is not None
    u = ctypes.windll.user32
    u.GetWindowLongPtrW.restype = ctypes.c_ssize_t
    assert u.GetWindowLongPtrW(ctypes.c_void_p(hook.hwnd), -20) & 0x10    # WS_EX_ACCEPTFILES


def test_selftest():
    assert main.selftest() == 0
