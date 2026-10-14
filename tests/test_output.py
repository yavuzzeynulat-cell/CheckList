import datetime as dt

import openpyxl

import compare
import detect
import output
from fixtures import CHECK_HEADER, RIA_HEADER, check, ria, xlsx

A = "26437-RIA-04C-QC-PIT-OP-00004"
D = "26437-RIA-02B-ST-IN2-RW-00571"


def test_two_sheets_full_rows(tmp_path):
    rr = ria(A)
    rr[4] = dt.datetime(2026, 10, 1)
    r = [detect.read(xlsx(tmp_path / "r.xlsx", RIA_HEADER, [rr, ria(D)]))]
    c = [detect.read(xlsx(tmp_path / "c.xlsx", CHECK_HEADER, [check("CHL-1", A)]))]
    p = output.write(compare.run(r, c), tmp_path / "out" / "CheckList Result.xlsx")
    wb = openpyxl.load_workbook(p)
    assert wb.sheetnames == [output.MATCHED_SHEET, output.MISSING_SHEET]
    m, x = wb[output.MATCHED_SHEET], wb[output.MISSING_SHEET]
    assert [c.value for c in m[1]] == RIA_HEADER + CHECK_HEADER
    assert m.cell(2, 2).value == A and m.cell(2, len(RIA_HEADER) + 1).value == "CHL-1"
    assert m.cell(2, 5).value == dt.datetime(2026, 10, 1)
    assert [c.value for c in x[1]] == RIA_HEADER
    assert x.cell(2, 2).value == D and x.max_row == 2
    assert m.freeze_panes == "A2" and m.auto_filter.ref


def test_unique_path(tmp_path):
    p = tmp_path / "a.xlsx"
    p.write_text("x")
    assert output.unique_path(p).name == "a (2).xlsx"


def test_out_dir_from_settings(tmp_path):
    import settings
    settings.save_out_dir(tmp_path)
    assert output.day_dir(dt.date(2026, 10, 14)) == tmp_path / "2026-10-14"
