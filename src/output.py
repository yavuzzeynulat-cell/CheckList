"""Sonuc Excel'i (GUI'siz): iki sayfa, satirlar tum sutunlariyla.

'Has Checklist': RIA'nin butun sutunlari + yaninda checklist'in butun
sutunlari; bir RIA'nin birden fazla checklist'i varsa her biri ayri satir.
'Missing Checklist': checklist'i olmayan RIA'lar, butun sutunlari.
(kullanici 2026-10-14: "full satiri verecek", "olmayanlarin da olanlarin da")
"""
import ctypes
import datetime as dt
import os
import tempfile
import uuid
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MATCHED_SHEET = "Has Checklist"
MISSING_SHEET = "Missing Checklist"
FILE_NAME = "CheckList Result.xlsx"
RIA_FILL = "DDEBF7"          # RIA sutunlarinin basligi
CHECK_FILL = "E2EFDA"        # checklist sutunlarinin basligi
MISSING_TAB = "F4B183"
MATCHED_TAB = "A9D08E"
MAX_WIDTH = 60
_DESKTOP = "{B4BFCC3A-DB2C-424C-B029-7FE99A87C641}"


class OutputError(Exception):
    pass


def desktop_dir():
    """Gercek masaustu (OneDrive yonlendirmesi dahil)."""
    try:
        class GUID(ctypes.Structure):
            _fields_ = [("d1", ctypes.c_ulong), ("d2", ctypes.c_ushort),
                        ("d3", ctypes.c_ushort), ("d4", ctypes.c_ubyte * 8)]
        u = uuid.UUID(_DESKTOP)
        g = GUID(u.fields[0], u.fields[1], u.fields[2], (ctypes.c_ubyte * 8).from_buffer_copy(u.bytes[8:]))
        out = ctypes.c_wchar_p()
        if ctypes.windll.shell32.SHGetKnownFolderPath(ctypes.byref(g), 0, None, ctypes.byref(out)) == 0:
            path = out.value
            ctypes.windll.ole32.CoTaskMemFree(out)
            return Path(path)
    except Exception:
        pass
    return Path.home() / "Desktop"


def base_dir():
    import settings
    return settings.load_out_dir() or desktop_dir() / "CheckList"


def day_dir(today=None):
    return base_dir() / f"{today or dt.date.today():%Y-%m-%d}"


def unique_path(path):
    path = Path(path)
    if not path.exists():
        return path
    n = 2
    while True:
        cand = path.with_name(f"{path.stem} ({n}){path.suffix}")
        if not cand.exists():
            return cand
        n += 1


def _width(v):
    if v is None:
        return 0
    if isinstance(v, (dt.datetime, dt.date)):
        return 16
    return max(len(line) for line in str(v).splitlines() or [""])


def _sheet(ws, header, fills, rows, tab):
    ws.sheet_properties.tabColor = tab
    thin = Side(style="thin", color="BFBFBF")
    ws.append(header)
    for j, fill in enumerate(fills, 1):
        c = ws.cell(1, j)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor=fill)
        c.border = Border(bottom=thin)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    widths = [len(h) for h in header]
    for r in rows:
        ws.append(r)
        for j, v in enumerate(r):
            widths[j] = max(widths[j], _width(v))
            if isinstance(v, dt.datetime):
                c = ws.cell(ws.max_row, j + 1)
                c.number_format = "dd/mm/yyyy" if (v.hour, v.minute) == (0, 0) else "dd/mm/yyyy hh:mm"
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = min(MAX_WIDTH, max(8, w + 2))
    ws.freeze_panes = "A2"
    if header:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{max(1, len(rows) + 1)}"


def write(result, out_path):
    out_path = Path(out_path)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = MATCHED_SHEET
    rh, ch = result.ria_header, result.check_header
    _sheet(ws, rh + ch, [RIA_FILL] * len(rh) + [CHECK_FILL] * len(ch),
           [list(r) + list(c) for r, c in result.matched], MATCHED_TAB)
    _sheet(wb.create_sheet(MISSING_SHEET), rh, [RIA_FILL] * len(rh), [list(r) for r in result.missing], MISSING_TAB)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=out_path.parent)
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, out_path)
    except OSError as e:
        if os.path.exists(tmp):
            os.remove(tmp)
        raise OutputError(f"Could not save {out_path.name} (open in Excel?). Close it and try again.") from e
    return out_path
