"""Birakilan Excel'in turunu basliktan tanir ve satirlarini okur (GUI'siz).

Olculdu (2026-10-14): checklist DataGrid'inde 'Task Number' + 'Related RIA',
RIA DataGrid'inde 'Document Number' + 'Document Title' var. Lab raporu
DataGrid'inde de 'Related RIA' var ama 'Task Number' yok: checklist sayilmaz.
Uc dosyanin adi da 'DataGrid ...' -- addan ayrilmaz, yalniz basliktan.
"""
import os
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

RIA, CHECKLIST, OTHER = "ria", "checklist", "other"
RIA_MARK = ("Document Number", "Document Title")
CHECKLIST_MARK = ("Task Number", "Related RIA")
RIA_KEY = "Document Number"
CHECKLIST_KEY = "Related RIA"
SOURCE_SHEET = "Sheet"
EXTENSIONS = (".xlsx", ".xlsm")


@dataclass
class Table:
    path: Path
    kind: str
    header: list = field(default_factory=list)   # tekrar olabilir: checklist'te iki 'Task Number'
    rows: list = field(default_factory=list)     # her satir header uzunlugunda
    error: str = ""


def kind_of(header):
    """(tur, hata). Hata: kullaniciya gosterilecek kisa neden."""
    names = set(header)
    if all(c in names for c in CHECKLIST_MARK):
        return CHECKLIST, ""
    if CHECKLIST_KEY in names:          # lab listesinde Document Number/Title da var: once bu
        return OTHER, "Lab report list, not a checklist"
    if all(c in names for c in RIA_MARK):
        return RIA, ""
    if RIA_KEY in names:
        return OTHER, "Has Document Number but no Document Title"
    return OTHER, "Not a RIA list or checklist"


def _clean(v):
    """Excel'in gosterdigi gibi: EDMS sondaki bosluklari atmadan yaziyor."""
    return v.strip() if isinstance(v, str) else v


def _blank(values):
    return all(v is None or (isinstance(v, str) and not v.strip()) for v in values)


def read(path):
    path = Path(path)
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except PermissionError:
        return Table(path, OTHER, error="Open in Excel — close it and drop again")
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError) as e:
        return Table(path, OTHER, error=f"Cannot read ({e.__class__.__name__})")
    try:
        ws = wb[SOURCE_SHEET] if SOURCE_SHEET in wb.sheetnames else wb.worksheets[0]
        if hasattr(ws, "reset_dimensions"):
            # EDMS <dimension ref="A1:C1"> yaziyor; read_only buna guvenip
            # 3 sutun / 1 satir goruyor (Handover'da olculdu).
            ws.reset_dimensions()
        it = ws.iter_rows(values_only=True)
        header = [str(v).strip() if v is not None else "" for v in next(it, ())]
        while header and not header[-1]:
            header.pop()
        kind, error = kind_of(header)
        rows = []
        if kind != OTHER:
            n = len(header)
            for r in it:
                values = [_clean(v) for v in r[:n]] + [None] * (n - len(r))
                if not _blank(values):
                    rows.append(values)
        return Table(path, kind, header, rows, error)
    finally:
        wb.close()


def _is_lock(p):
    return p.name.startswith("~$")


def expand_drop(paths):
    """Birakilan yollar -> Excel dosyalari. Klasor birakilirsa icindekiler."""
    out = []
    for raw in paths:
        p = Path(raw)
        if p.is_dir():
            out += sorted(f for f in p.iterdir()
                          if f.is_file() and f.suffix.lower() in EXTENSIONS and not _is_lock(f))
        elif p.is_file() and p.suffix.lower() in EXTENSIONS and not _is_lock(p):
            out.append(p)
    return out


def norm(p):
    return os.path.normcase(os.path.abspath(str(p)))
